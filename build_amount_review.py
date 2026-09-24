"""
build_amount_review.py — writes amount_review.xlsx, a queue of the records
whose amounts most deserve a human check against the institution's own page.

Run:
    python build_amount_review.py            # refuses to overwrite
    python build_amount_review.py --force    # rebuild, DISCARDING decisions

Why a manual queue: both of this project's large amount errors (IDB Invest's
group-level figure, AfDB's programme envelope) were found by accident, from a
chart that looked wrong. The 200 largest records hold about an eighth of all
dollars, so checking them by hand is the cheapest insurance there is.

What goes in, in this order:
  A  suspects: records that match a known error pattern (an amount shared
     with other records of the same institution, a trade-finance line, or an
     IFC 'World Region' record that kept a programme envelope);
  B  the rest of the 200 largest records by USD;
  C  each institution's five largest, so the small ones get looked at too.

The sheet is the reviewer's working file and holds their decisions, so this
script will not overwrite it without --force. Decisions are read back and
applied separately; nothing here changes the database, which is opened
read-only.

Matching a decision back to its record: never by projects.id (reassigned on
every reload). Use institution + source URL + approval date + exact name +
source amount ROUNDED TO THE CENT + currency. Name and URL alone are not
enough (EBRD's URL is one spreadsheet for all 9,415 records; DFC has 29
'Redacted'); the amount must be rounded because Excel keeps 15 significant
digits. Verified unique for all 215 rows on 2026-09-24.
"""

import re
import sqlite3
import sys
from datetime import date
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
DB_PATH = HERE / "data" / "dfi_tracker.db"
OUT_PATH = HERE / "amount_review.xlsx"

TOP_N = 200
PER_INSTITUTION = 5

VERDICTS = [
    "Correct - matches the page",
    "Correct - small difference (rounding or exchange rate)",
    "Wrong - programme or envelope total, not this deal",
    "Wrong - total project cost, not this institution's share",
    "Wrong - cumulative or revolving total, not one commitment",
    "Wrong - currency or units",
    "Wrong - other (explain in Notes)",
    "Can't verify - no link, page gone, or no amount shown",
    "Unsure - discuss",
]

TRADE_FINANCE = re.compile(
    r"\b(TFP|TFFP|GTFP|GTLP|GTSF|GSCF|TECF)\b|trade financ|supply chain financ",
    re.IGNORECASE)

# Where the amount sits on each institution's page, and what difference is
# EXPECTED rather than an error. "Checked" is whether the page layout was
# looked at while building this sheet (2026-09-24) — unchecked ones are
# the reviewer's to confirm.
GUIDE = {
    "IFC": ("Open the link. Click the section 'Total Project Cost and Amount "
            "and Nature of IFC's Investment'. It gives the total project and "
            "IFC's own investment separately; compare OUR amount with IFC's "
            "investment.",
            "Loan plus equity may be listed separately; add them. 'Up to' "
            "figures are fine.", "Yes"),
    "EBRD": ("There is no page per record: the link is a web search. Open the "
             "result whose address contains '/projects/psd/'. That is the "
             "Project Summary Document; its financing section states EBRD's "
             "own amount.",
             "Our figure is EBRD's 'Net Cumulative Bank Investment' from its "
             "portfolio list, not the signing amount, so it can differ from "
             "the PSD. A much LARGER figure than the PSD is exactly what to "
             "look for.", "Search only; PSD layout not checked"),
    "EIB Global": ("Open the link and click 'Signature(s)'. It lists the "
                   "signed amount and each signature date.",
                   "EIB publishes loan tranches, so one project can be "
                   "several of our records. Match OUR amount to one "
                   "signature line, not the project total.", "Yes"),
    "AfDB": ("Open the link and read 'Commitment (UA)' on Overview. Then "
             "click 'Financial information': under 'Commitments by source' "
             "note the African Development Bank Group line.",
             "Amounts are in UA (Units of Account, the IMF's SDR). Record UA "
             "figures, not USD. If 'Commitments by source' shows other "
             "financiers, OUR figure may include them — say so in Notes.",
             "Yes"),
    "IDB Invest": ("Open the link. 'Investment summary' shows 'FINANCING "
                   "AMOUNT'.",
                   "The page can show the whole IDB Group's figure, typically "
                   "1.8 to 5.3 times IDB Invest's own share, which is what we "
                   "load. A LARGER page figure is expected; a SMALLER one is "
                   "a problem.", "Yes"),
    "FMO": ("Open the link. Read 'Total FMO financing' in the side panel.",
            "'Funding' says whose money it was: 'FMO NV' is FMO's own "
            "account; anything else is a government fund FMO manages.",
            "Yes"),
    "DFC": ("The link is a PDF project summary. Look for the proposed DFC "
            "financing amount.",
            "Records named 'Redacted' have no public page: mark 'Can't "
            "verify'.", "No"),
    "ADB": ("Open the link (your browser may show a security check first). "
            "Look for the financing table, ADB's own line.",
            "One ADB record is in Indian rupees; its description says 'up to "
            "$750 million'.", "No - the site blocked automated checking"),
    "BII": ("The link opens d-portal (the IATI data viewer). Look at the "
            "commitment transactions for this activity.",
            "BII publishes transactions, so one deal can be several records.",
            "No"),
    "Proparco": ("Open the link (French). Look for 'Montant' / the amount "
                 "of Proparco's financing. Records loaded from AFD's open data "
                 "have no page of their own: the link is a web search instead.",
                 "Amounts are in EUR.", "No"),
}

# Colours: the house palette's deep blue for headers; yellow marks the cells
# the reviewer fills in; grey marks what comes from the database.
HEADER_FILL = PatternFill("solid", fgColor="0E2A3F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2B3")
LOCKED_FILL = PatternFill("solid", fgColor="EEF1F3")
EXAMPLE_FILL = PatternFill("solid", fgColor="E2EFDA")
FONT = "Arial"
THIN = Side(style="thin", color="C9D1D7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fetch(conn):
    """The queue, in review order, with the context each row needs."""
    rows = conn.execute(
        """SELECT id, institution, project_name, canonical_country, country,
                  COALESCE(CAST(strftime('%Y', approval_date) AS INTEGER),
                           fiscal_year) AS year,
                  approval_date, amount_original, currency, amount_usd,
                  source_url, description
           FROM projects WHERE amount_usd IS NOT NULL AND amount_usd > 0
           ORDER BY amount_usd DESC""").fetchall()
    cols = ["id", "institution", "name", "country", "raw_country", "year",
            "approval_date", "amount_original", "currency", "amount_usd",
            "url", "description"]
    recs = [dict(zip(cols, r)) for r in rows]

    # Rank within institution, and how many OTHER records of the same
    # institution carry exactly this amount on the same approval date: an
    # envelope stamped onto each deal looks like this. The date matters —
    # without it, round figures (EUR 500m, USD 400m) match by coincidence.
    def key(r):
        return (r["institution"], r["amount_original"], r["currency"],
                r["approval_date"])
    same = {}
    for r in recs:
        if r["approval_date"]:
            same[key(r)] = same.get(key(r), 0) + 1
    rank = {}
    for r in recs:
        rank[r["institution"]] = rank.get(r["institution"], 0) + 1
        r["inst_rank"] = rank[r["institution"]]
        r["same_amount"] = same.get(key(r), 1) - 1

    top_ids = {r["id"] for r in recs[:TOP_N]}
    chosen = {r["id"]: r for r in recs[:TOP_N]}
    for r in recs:
        if r["inst_rank"] <= PER_INSTITUTION and r["id"] not in chosen:
            chosen[r["id"]] = r

    for r in chosen.values():
        r["reasons"] = reasons(r)
        r["priority"] = ("A" if r["reasons"] else
                         "B" if r["id"] in top_ids else "C")
        r["context"] = context(r)
    return sorted(chosen.values(),
                  key=lambda r: ("ABC".index(r["priority"]), -r["amount_usd"]))


def reasons(r):
    """Why a record is a SUSPECT. Empty list = only here for its size."""
    out = []
    if r["same_amount"] >= 1 and r["amount_usd"] >= 100e6:
        out.append(f"{r['same_amount']} other {r['institution']} record(s) "
                   "carry exactly this amount on the same date")
    if TRADE_FINANCE.search(r["name"] or ""):
        out.append("trade-finance line")
    if (r["institution"] == "IFC" and (r["raw_country"] or "").strip().lower()
            == "world region" and r["amount_usd"] >= 500e6):
        out.append("IFC 'World Region' record of USD 500m or more")
    return out


def context(r):
    """What we already know, so the reviewer doesn't rediscover it."""
    notes = []
    if "IFC 'World Region'" in " ".join(r["reasons"]):
        notes.append(
            "IFC's trade programmes stamp their whole envelope on partner "
            "records. Since 24 Sep 2026 the loader blanks it wherever several "
            "World Region records share one (GSCF, GTSF), but a LONE World "
            "Region record keeps its amount as the programme's own. Check "
            "whether this is the programme itself or one partner bank.")
    if "trade-finance line" in r["reasons"]:
        notes.append(
            "Trade lines revolve: check whether the figure is a LIMIT (one "
            "commitment) or a running total of trade covered over the years.")
    if r["institution"] == "EBRD":
        notes.append("Our figure is EBRD's 'Net Cumulative Bank Investment'.")
    if r["institution"] == "DFC" and (r["name"] or "") == "Redacted":
        notes.append("DFC withheld the name; there is no public page.")
    if r["currency"] == "INR":
        notes.append("Loaded in rupees; the description says 'up to $750 "
                     "million'.")
    return " ".join(notes)


def link(r):
    """(url, label) for the source. EBRD's records, and Proparco's loaded from
    AFD open data, link to a whole dataset rather than a page per deal, so
    they get a web search for the deal instead."""
    url = r["url"] or ""
    if r["institution"] == "EBRD":
        q = quote_plus(f'site:ebrd.com/home/work-with-us/projects/psd {r["name"]}')
        return f"https://www.google.com/search?q={q}", "Search EBRD"
    if r["institution"] == "Proparco" and "opendata.afd.fr" in url:
        q = quote_plus(f'site:proparco.fr {r["name"]}')
        return f"https://www.google.com/search?q={q}", "Search Proparco"
    if not url.startswith("http"):
        return None, "No link"
    return url, "Open page"


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def write_start_here(wb, n_rows, counts, built_from):
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    lines = [
        ("title", "Amount review — checking our largest records against the source"),
        ("small", f"Built {date.today():%d %B %Y} from the database as loaded "
                  f"{built_from}. {n_rows} records: {counts['A']} suspects (A), "
                  f"{counts['B']} large (B), {counts['C']} institution top-fives (C)."),
        ("h", "What this is for"),
        ("p", "Our two biggest amount errors so far (IDB Invest's group figure, AfDB's "
              "programme envelope) were found by accident. This queue checks the "
              "records that carry the most dollars, so the next one is found on "
              "purpose. The 200 largest records alone are about an eighth of every "
              "dollar in the tracker."),
        ("p", f"Each check takes one to three minutes. At five a day, priority A "
              f"takes about {-(-counts['A'] // 5)} days and the whole queue about "
              f"{-(-n_rows // 5)}. Priority A is where the errors should be; the "
              f"rest is confirmation."),
        ("h", "Step by step"),
        ("n", "1.  Open the 'Queue' tab. Work top to bottom: it is already sorted, "
              "suspects first, then by size."),
        ("n", "2.  Find the first row with an empty yellow 'Verdict' cell."),
        ("n", "3.  Read the grey cells: who, what, OUR amount in the source's own "
              "currency, and the 'Already known' note if there is one."),
        ("n", "4.  Click the link in 'Source' ('Open page', or a web search for EBRD and "
              "some Proparco records). Use the 'Institution guide' "
              "tab to find where that institution puts the amount, and which "
              "differences are normal for it."),
        ("n", "5.  Find the institution's OWN commitment for this deal. Not the total "
              "project cost, not a programme or facility total, not co-financiers' "
              "money."),
        ("n", "6.  Fill the yellow cells: pick a Verdict from the dropdown; type the "
              "amount the page shows, in the page's currency, as a plain number "
              "(250000000, not '250m'); the currency; the words the page uses for it; "
              "and today's date (Ctrl+; types it)."),
        ("n", "7.  If the verdict starts 'Wrong', also say in Notes where on the page "
              "you found the right figure. The page amount cell turns red until you "
              "fill it in."),
        ("n", "8.  Save. That's all: nothing changes in the database until you ask "
              "Claude to apply the review ('apply the amount review')."),
        ("h", "Choosing a verdict"),
        ("p", "Correct - matches the page: same figure, or within 1%."),
        ("p", "Correct - small difference: within about 5%, explained by rounding or "
              "the exchange rate on a different date."),
        ("p", "Wrong - programme or envelope total: the figure belongs to a whole "
              "programme or facility (the IFC GSCF case below)."),
        ("p", "Wrong - total project cost: the figure is the whole project, including "
              "other lenders and the sponsor."),
        ("p", "Wrong - cumulative or revolving total: typically trade lines, where "
              "the page gives a limit and our figure is much larger."),
        ("p", "Wrong - currency or units: right digits, wrong currency, or out by a "
              "factor of 1,000 or 1,000,000."),
        ("p", "Can't verify: no link, a dead page, or a page with no amount. Say which "
              "in Notes. This is a legitimate result, not a failure."),
        ("p", "Unsure - discuss: anything you'd rather talk through. Say why in Notes."),
        ("h", "Worked example (a real one, found while building this sheet)"),
        ("p", "IFC 'GSCF Citi II', 2022. Our amount was USD 3,115,000,000. The page, "
              "under 'Total Project Cost and Amount and Nature of IFC's Investment', "
              "says: 'The IFC Investment will be in amount of up to US$250 million'. "
              "This one led to a loader fix on 24 September 2026, so it is no longer "
              "in the queue, but it is exactly what an error looks like. Filled in, "
              "it would read:"),
        ("ex", "Verdict: Wrong - programme or envelope total, not this deal  |  Amount "
               "on page: 250000000  |  Currency: USD  |  Page calls it: IFC "
               "Investment (up to)  |  Notes: USD 3,115m is the GSCF programme "
               "figure; GSCF-SMBC carries the same number."),
        ("h", "Rules"),
        ("p", "Only type in the yellow columns. The grey ones come from the database, "
              "and the columns at the far right are how your decision is matched back "
              "to the record."),
        ("p", "Sorting and filtering are fine: every row carries its own ID."),
        ("p", "Don't rebuild this file: 'python build_amount_review.py' refuses to "
              "overwrite it, because it holds your decisions."),
        ("p", "The 'Progress' tab counts what you've done."),
    ]
    row = 2
    for kind, text in lines:
        if kind == "h":
            row += 1  # a blank line above each heading
        cell = ws.cell(row=row, column=2, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(name=FONT, bold=True, size=16, color="0E2A3F")
        elif kind == "small":
            cell.font = Font(name=FONT, size=9, color="5D6C77")
        elif kind == "h":
            cell.font = Font(name=FONT, bold=True, size=12, color="0E2A3F")
        elif kind == "ex":
            cell.font = Font(name=FONT, size=10)
            cell.fill = EXAMPLE_FILL
        else:
            cell.font = Font(name=FONT, size=10)
        # Rough height so wrapped lines show without manual resizing.
        ws.row_dimensions[row].height = max(15, 15 * (len(text) // 105 + 1))
        row += 1


def write_queue(wb, recs):
    ws = wb.create_sheet("Queue")
    headers = [
        # (header, width, kind) — kind: db (grey), in (yellow), key (grey, far right)
        ("#", 5, "db"), ("Review ID", 9, "db"), ("Priority", 8, "db"),
        ("Why it's here", 30, "db"), ("Institution", 11, "db"),
        ("Project name", 38, "db"), ("Country", 16, "db"), ("Year", 7, "db"),
        ("Our amount (source currency)", 18, "db"), ("Currency", 9, "db"),
        ("Our amount (USD m)", 11, "db"),
        ("Rank within institution", 10, "db"),
        ("Already known", 45, "db"), ("Source", 12, "db"),
        ("Verdict", 34, "in"), ("Amount on page", 17, "in"),
        ("Page currency", 9, "in"), ("What the page calls it", 24, "in"),
        ("Notes", 40, "in"), ("Date reviewed", 12, "in"),
        ("Source URL (key)", 30, "key"), ("Approval date (key)", 12, "key"),
        ("Exact name (key)", 30, "key"),
    ]
    for i, (h, w, _) in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws, 1, len(headers))
    for i, (_, _, kind) in enumerate(headers, 1):
        if kind == "in":
            ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="A8853F")
    ws.row_dimensions[1].height = 42

    for n, r in enumerate(recs, 1):
        row = n + 1
        url, label = link(r)
        why = "; ".join(r["reasons"]) if r["reasons"] else (
            f"top {TOP_N} by USD" if r["priority"] == "B"
            else f"{r['institution']}'s top {PER_INSTITUTION}")
        values = [n, f"AR-{n:03d}", r["priority"], why, r["institution"],
                  r["name"], r["country"], r["year"], r["amount_original"],
                  r["currency"], r["amount_usd"] / 1e6, r["inst_rank"],
                  r["context"], None, None, None, None, None, None, None,
                  r["url"], r["approval_date"], r["name"]]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            kind = headers[c - 1][2]
            cell.font = Font(name=FONT, size=9)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=kind != "key")
            cell.fill = INPUT_FILL if kind == "in" else LOCKED_FILL
        src = ws.cell(row=row, column=14)
        if url:
            src.value = f'=HYPERLINK("{url}","{label}")'
            src.font = Font(name=FONT, size=9, color="2E6187", underline="single")
        else:
            src.value = label
        ws.cell(row=row, column=9).number_format = "#,##0"
        ws.cell(row=row, column=11).number_format = "#,##0.0"
        ws.cell(row=row, column=16).number_format = "#,##0"
        ws.cell(row=row, column=20).number_format = "yyyy-mm-dd"

    last = len(recs) + 1
    dv = DataValidation(type="list", formula1="=Lists!$A$1:$A$%d" % len(VERDICTS),
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Pick from the list",
                        error="Choose a verdict from the dropdown.")
    ws.add_data_validation(dv)
    dv.add(f"O2:O{last}")
    num = DataValidation(type="decimal", operator="greaterThanOrEqual",
                         formula1="0", allow_blank=True, showErrorMessage=True,
                         error="Type a plain number, e.g. 250000000.")
    ws.add_data_validation(num)
    num.add(f"P2:P{last}")

    red = PatternFill("solid", fgColor="F4B6B0")
    ws.conditional_formatting.add(
        f"P2:P{last}", FormulaRule(formula=['AND(LEFT($O2,5)="Wrong",$P2="")'],
                                   fill=red))
    ws.conditional_formatting.add(
        f"A2:C{last}", FormulaRule(formula=['LEFT($O2,7)="Correct"'],
                                   fill=PatternFill("solid", fgColor="CFE3D0")))
    ws.conditional_formatting.add(
        f"A2:C{last}", FormulaRule(formula=['LEFT($O2,5)="Wrong"'], fill=red))
    ws.conditional_formatting.add(
        f"A2:C{last}", FormulaRule(formula=['OR(LEFT($O2,5)="Can\'t",'
                                            'LEFT($O2,6)="Unsure")'],
                                   fill=PatternFill("solid", fgColor="F2DDB0")))
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    return last


def write_guide(wb):
    ws = wb.create_sheet("Institution guide")
    headers = [("Institution", 14), ("Where the amount is", 60),
               ("Differences that are NORMAL here", 60),
               ("Page layout checked 24 Sep 2026?", 18)]
    for i, (h, w) in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws, 1, len(headers))
    for n, (inst, (where, normal, checked)) in enumerate(GUIDE.items(), 2):
        for c, v in enumerate([inst, where, normal, checked], 1):
            cell = ws.cell(row=n, column=c, value=v)
            cell.font = Font(name=FONT, size=10, bold=c == 1)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
    ws.freeze_panes = "B2"


def write_progress(wb, last):
    ws = wb.create_sheet("Progress")
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 10
    ws["A1"] = "Progress"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="0E2A3F")
    for i, h in enumerate(["", "All", "A", "B", "C"], 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, 5)
    q = "Queue"
    rng_v, rng_p = f"{q}!$O$2:$O${last}", f"{q}!$C$2:$C${last}"
    rows = [("Records in the queue", None, "total"),
            ("Reviewed (any verdict)", None, "done")]
    rows += [(v, v, "verdict") for v in VERDICTS]
    rows += [("Still to do", None, "todo")]
    r = 4
    for label, verdict, kind in rows:
        ws.cell(row=r, column=1, value=label)
        for c, pr in enumerate([None, "A", "B", "C"], 2):
            if kind == "total":
                f = (f"=COUNTA({q}!$B$2:$B${last})" if pr is None else
                     f'=COUNTIF({rng_p},"{pr}")')
            elif kind == "done":
                f = (f"=COUNTA({rng_v})" if pr is None else
                     f'=COUNTIFS({rng_p},"{pr}",{rng_v},"<>")')
            elif kind == "todo":
                f = f"={get_column_letter(c)}4-{get_column_letter(c)}5"
            else:
                f = (f'=COUNTIF({rng_v},$A{r})' if pr is None else
                     f'=COUNTIFS({rng_p},"{pr}",{rng_v},$A{r})')
            ws.cell(row=r, column=c, value=f)
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, bold=kind in ("total", "todo"))
            cell.border = BORDER
        r += 1
    ws.cell(row=r + 1, column=1,
            value="Counts update as you fill the Verdict column. 'Reviewed' counts "
                  "any row with a verdict.").font = Font(name=FONT, size=9,
                                                         color="5D6C77")


def write_lists(wb):
    ws = wb.create_sheet("Lists")
    for i, v in enumerate(VERDICTS, 1):
        ws.cell(row=i, column=1, value=v)
    ws.sheet_state = "hidden"


def main():
    if OUT_PATH.exists() and "--force" not in sys.argv:
        sys.exit(f"{OUT_PATH.name} already exists and may hold review decisions. "
                 "Rerun with --force to rebuild it and DISCARD them.")
    conn = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    try:
        recs = fetch(conn)
        built_from = conn.execute(
            "SELECT MAX(scraped_at) FROM projects").fetchone()[0][:10]
    finally:
        conn.close()
    counts = {p: sum(r["priority"] == p for r in recs) for p in "ABC"}

    wb = Workbook()
    write_start_here(wb, len(recs), counts, built_from)
    last = write_queue(wb, recs)
    write_guide(wb)
    write_progress(wb, last)
    write_lists(wb)
    wb.active = 0
    # openpyxl stores formulas without results; make Excel compute them all
    # the moment the file opens, or the Progress tab shows blanks.
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT_PATH)
    print(f"Wrote {len(recs)} records to {OUT_PATH.name}: "
          f"A {counts['A']}, B {counts['B']}, C {counts['C']}")


if __name__ == "__main__":
    main()
